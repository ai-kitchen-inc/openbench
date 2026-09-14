"""Tests for General Chat custom skill store and admin routes."""

from __future__ import annotations

import sys
import tempfile
import unittest
from contextlib import ExitStack
from os import environ
from pathlib import Path
from unittest.mock import Mock, patch

from fastapi.testclient import TestClient

from openbench.integrations.firebase_auth import FirebaseUser

GENERAL_CHAT_SRC = Path(__file__).resolve().parents[1] / "examples" / "general-chat" / "src"
if str(GENERAL_CHAT_SRC) not in sys.path:
    sys.path.insert(0, str(GENERAL_CHAT_SRC))

from general_chat.server.custom_functions import CustomFunctionStore  # noqa: E402
from general_chat.server.custom_skills import CustomSkillError, CustomSkillStore  # noqa: E402

ADMIN = FirebaseUser(uid="admin-1", email="boss@example.com")
ADMIN_H = {"Authorization": "Bearer token-admin"}


class TestCustomSkillStore(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        env_patch = patch.dict(
            environ,
            {"GENERAL_CHAT_CUSTOM_SKILLS_DIR": "", "CUSTOM_FN_DATA_PATH": ""},
        )
        env_patch.start()
        self.addCleanup(env_patch.stop)
        self.store = CustomSkillStore(self._tmp.name)

    def test_save_list_delete_round_trip(self):
        saved = self.store.save(
            "risk-review",
            name="Risk Review",
            description="Reviews decision risks.",
            triggers=["risk", "mitigation"],
            instructions="Always return risks, impact, and mitigation.",
        )
        self.assertEqual(saved["id"], "risk-review")
        self.assertEqual(saved["name"], "Risk Review")
        self.assertIn("risk", saved["triggers"])
        self.assertIn("Always return", saved["skill_md"])
        paths = self.store.paths()
        self.assertEqual(len(paths), 1)
        self.assertTrue((paths[0] / "SKILL.md").is_file())

        listed = self.store.list()
        self.assertEqual([skill["id"] for skill in listed], ["risk-review"])
        self.assertTrue(self.store.delete("risk-review"))
        self.assertEqual(self.store.list(), [])
        self.assertFalse(self.store.delete("risk-review"))

    def test_save_rejects_invalid_payload(self):
        with self.assertRaises(CustomSkillError):
            self.store.save("1-bad", name="Bad", instructions="Do things.")
        with self.assertRaises(CustomSkillError):
            self.store.save("ok-skill", name="", instructions="Do things.")
        with self.assertRaises(CustomSkillError):
            self.store.save("ok-skill", name="OK", instructions="")
        with self.assertRaises(CustomSkillError):
            self.store.save("ok-skill", name="OK", instructions="Do things.", version="vNext")

    def test_save_from_prompt_generates_unique_markdown(self):
        first = self.store.save_from_prompt(
            "Buat skill untuk review kontrak vendor. Agent harus menilai klausul risiko, "
            "kewajiban, dan rekomendasi negosiasi."
        )
        second = self.store.save_from_prompt(
            "Buat skill untuk review kontrak vendor. Agent harus menilai klausul risiko."
        )
        self.assertNotEqual(first["id"], second["id"])
        self.assertTrue(first["id"].startswith("review-kontrak-vendor"))
        self.assertIn("## Instructions", first["skill_md"])
        self.assertIn("## Triggers", first["skill_md"])

    def test_save_from_prompt_reuses_existing_custom_function(self):
        functions = CustomFunctionStore(self._tmp.name)
        functions.save(
            "custom_skill_image_brief",
            "\n".join(
                [
                    "def custom_skill_image_brief(project_description):",
                    "    return {",
                    "        'computed_by': 'custom_skill_image_brief',",
                    "        'prompt': project_description,",
                    "    }",
                ]
            ),
            "Existing image brief function.",
        )

        saved = self.store.save_from_prompt(
            "Buat skill rancangan proyek pembangunan yang juga butuh gambar visual.",
            custom_functions=functions,
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"], [])
        self.assertEqual(tooling["required"][0]["name"], "custom_skill_image_brief")
        self.assertEqual(tooling["required"][0]["type"], "custom_function")
        self.assertIn("## Tooling", saved["skill_md"])
        self.assertIn("custom_function_run_function", saved["skill_md"])

    def test_save_from_prompt_reuses_manual_function_by_semantic_name(self):
        functions = CustomFunctionStore(self._tmp.name)
        functions.save(
            "add",
            "\n".join(
                [
                    "def add(a, b):",
                    "    return {'computed_by': 'add', 'result': float(a) + float(b)}",
                ]
            ),
            "Menjumlahkan dua angka.",
        )

        saved = self.store.save_from_prompt(
            "Buat skill untuk menghitung penjumlahan dua angka. Saat user memberi "
            "dua angka, hasil harus akurat dan konsisten.",
            custom_functions=functions,
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"], [])
        self.assertEqual(tooling["required"][0]["name"], "add")
        self.assertIn('name="add"', saved["skill_md"])

    def test_save_from_prompt_creates_missing_custom_function(self):
        functions = CustomFunctionStore(self._tmp.name)

        saved = self.store.save_from_prompt(
            "Buat skill estimasi anggaran proyek pembangunan dari bahan baku dan biaya.",
            custom_functions=functions,
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"][0]["name"], "custom_skill_estimate_budget")
        self.assertTrue((functions.root / "custom_skill_estimate_budget.py").is_file())
        self.assertIn("custom_skill_estimate_budget", saved["skill_md"])

    def test_save_from_prompt_recommends_function_for_plain_calculation_prompt(self):
        functions = CustomFunctionStore(self._tmp.name)

        saved = self.store.save_from_prompt(
            "Buat skill untuk menghitung luas bangunan persegi panjang. Saat user "
            "memberikan panjang dan lebar tanah, berikan hasil luas yang akurat "
            "dan konsisten.",
            custom_functions=functions,
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"][0]["name"], "hitung_luas_bangunan")
        self.assertTrue((functions.root / "hitung_luas_bangunan.py").is_file())
        self.assertIn('name="hitung_luas_bangunan"', saved["skill_md"])

    def test_save_from_prompt_creates_explicit_named_custom_function(self):
        functions = CustomFunctionStore(self._tmp.name)

        saved = self.store.save_from_prompt(
            "Saya ingin membuat skill menghitung luas bangunan berbentuk persegi panjang "
            "ketika user melampirkan panjang dan lebar tanah. Perhitungan harus membuat "
            "fungsi kustom dengan nama fungsi luas_bangunan.",
            custom_functions=functions,
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"][0]["name"], "luas_bangunan")
        self.assertTrue((functions.root / "luas_bangunan.py").is_file())
        code = (functions.root / "luas_bangunan.py").read_text(encoding="utf-8")
        self.assertIn("def luas_bangunan", code)
        self.assertIn("panjang_value * lebar_value", code)

    def test_save_from_prompt_keeps_plain_guidance_skill_knowledge_only(self):
        functions = CustomFunctionStore(self._tmp.name)

        saved = self.store.save_from_prompt(
            "Buat skill untuk membantu menulis balasan customer service dengan nada "
            "ramah, singkat, dan profesional.",
            custom_functions=functions,
        )

        self.assertEqual(saved["tooling"]["required"], [])
        self.assertEqual(functions.names(), set())

    def test_save_from_prompt_creates_prompt_references_when_template_is_described(self):
        saved = self.store.save_from_prompt(
            "Buat skill invoice analyzer dengan SOP validasi nomor invoice, pajak, dan total. "
            "Output harus mengikuti template laporan: Bab 1 Ringkasan, Bab 2 Temuan, "
            "Bab 3 Rekomendasi."
        )

        paths = self.store.paths()
        self.assertEqual(len(paths), 1)
        self.assertTrue((paths[0] / "references" / "prompt-rules.md").is_file())
        self.assertTrue((paths[0] / "references" / "output-template.md").is_file())
        self.assertIn("references/prompt-rules.md", saved["skill_md"])
        self.assertEqual(len(saved["resources"]["references"]), 2)

    def test_save_from_prompt_persists_uploaded_assets_and_text_references(self):
        upload_dir = Path(self._tmp.name) / "uploads"
        upload_dir.mkdir()
        sop = upload_dir / "invoice-sop.txt"
        sop.write_text("Validasi PPN dan nomor invoice.", encoding="utf-8")
        template = upload_dir / "invoice-template.xlsx"
        template.write_bytes(b"fake-xlsx")

        saved = self.store.save_from_prompt(
            "Buat skill invoice analyzer.",
            uploads=[
                {
                    "filename": "invoice-sop.txt",
                    "source_name": "invoice-sop.txt",
                    "bucket": "references",
                    "description": "SOP invoice.",
                    "content": "# Invoice SOP\n\nValidasi PPN dan nomor invoice.",
                },
                {
                    "filename": "invoice-template.xlsx",
                    "source_name": "invoice-template.xlsx",
                    "bucket": "assets",
                    "description": "Template laporan invoice.",
                    "source_path": str(template),
                },
            ],
        )

        skill_dir = self.store.paths()[0]
        self.assertTrue((skill_dir / "references" / "invoice-sop.md").is_file())
        self.assertTrue((skill_dir / "assets" / "invoice-template.xlsx").is_file())
        self.assertEqual(len(saved["resources"]["references"]), 1)
        self.assertEqual(len(saved["resources"]["assets"]), 1)

    def test_upload_classification_handles_sop_text_and_excel_template_together(self):
        upload_dir = Path(self._tmp.name) / "uploads"
        upload_dir.mkdir()
        sop = upload_dir / "sop-kategori-nilai.txt"
        sop.write_text(
            "Jika nilai >=75 maka Diatas KKM. Jika nilai <75 maka Dibawah KKM. "
            "Jika nilai '-' maka Tidak mengikuti ujian.",
            encoding="utf-8",
        )
        template = upload_dir / "template-nilai.xlsx"
        try:
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("openpyxl is required for Excel template extraction")
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Nilai Siswa"
        worksheet.append(["Nama", "Nilai (dalam angka)", "Nilai (dalam tulisan)", "Keterangan"])
        workbook.save(template)
        prompt = (
            "Saya ingin membuat skill untuk menginput data nilai siswa kedalam template "
            "file excel yang sudah saya berikan. Pemberian kategori nilai mengikuti SOP "
            "file yang sudah saya berikan juga."
        )

        saved = self.store.save_from_prompt(
            prompt,
            uploads=[
                {
                    "filename": "sop-kategori-nilai.txt",
                    "source_name": "sop-kategori-nilai.txt",
                    "bucket": "references",
                    "text_preview": sop.read_text(encoding="utf-8"),
                    "source_path": str(sop),
                },
                {
                    "filename": "template-nilai.xlsx",
                    "source_name": "template-nilai.xlsx",
                    "bucket": "assets",
                    "source_path": str(template),
                },
            ],
        )

        skill_dir = self.store.paths()[0]
        self.assertTrue((skill_dir / "references" / "sop-kategori-nilai.md").is_file())
        self.assertTrue((skill_dir / "assets" / "template-nilai.xlsx").is_file())
        sop_reference = (skill_dir / "references" / "sop-kategori-nilai.md").read_text(
            encoding="utf-8"
        )
        self.assertIn("Tidak mengikuti ujian", sop_reference)
        self.assertIn("assets/template-nilai.xlsx", saved["skill_md"])
        template_reference = skill_dir / "references" / "excel-template-template-nilai.md"
        self.assertTrue(template_reference.is_file())
        template_reference_text = template_reference.read_text(encoding="utf-8")
        self.assertIn("Nama", template_reference_text)
        self.assertIn("Nilai (dalam angka)", template_reference_text)
        self.assertIn("Nilai (dalam tulisan)", template_reference_text)
        self.assertIn("Keterangan", template_reference_text)
        self.assertTrue(saved["template_tool"].startswith("fill_"))

        from openbench.intelligence.skill import Skill

        loaded_skill = Skill.from_dir(skill_dir)
        tools = {name: fn for name, fn, _ in loaded_skill.tools}
        self.assertIn(saved["template_tool"], tools)
        result = tools[saved["template_tool"]](
            asset_path="assets/template-nilai.xlsx",
            records=[
                {
                    "Nama": "Andi",
                    "Nilai (dalam angka)": 60,
                    "Nilai (dalam tulisan)": "enam puluh",
                    "Keterangan": "Dibawah KKM",
                }
            ],
            output_filename="hasil-nilai.xlsx",
        )
        self.assertNotIn("error", result)
        self.assertEqual(
            result["filledColumns"],
            ["Nama", "Nilai (dalam angka)", "Nilai (dalam tulisan)", "Keterangan"],
        )

    def test_docx_template_tool_fills_form_table_and_refuses_blank_output(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is required for DOCX template verification")

        upload_dir = Path(self._tmp.name) / "uploads"
        upload_dir.mkdir()
        template = upload_dir / "form-penilaian-siswa.docx"
        document = Document()
        document.add_heading("Form Penilaian Siswa", level=1)
        table = document.add_table(rows=6, cols=2)
        rows = [
            ("Nama", ""),
            ("Kelas", ""),
            ("Fisika", ""),
            ("Kimia", ""),
            ("Matematika", ""),
            ("Biologi", ""),
        ]
        for row_index, (label, value) in enumerate(rows):
            table.cell(row_index, 0).text = label
            table.cell(row_index, 1).text = value
        document.save(template)

        saved = self.store.save_from_prompt(
            "Saya ingin membuat skill untuk menginput form penilaian siswa kedalam template "
            "file docs yang sudah saya berikan. Keluarkan output yang sesuai dengan template "
            "jangan ubah template sedikitpun.",
            uploads=[
                {
                    "filename": "form-penilaian-siswa.docx",
                    "source_name": "form-penilaian-siswa.docx",
                    "bucket": "assets",
                    "source_path": str(template),
                }
            ],
        )

        from openbench.intelligence.skill import Skill

        skill_dir = self.store.paths()[0]
        template_map = skill_dir / "references" / "document-template-form-penilaian-siswa.md"
        self.assertTrue(template_map.is_file())
        template_map_text = template_map.read_text(encoding="utf-8")
        self.assertIn("Suggested Field Keys", template_map_text)
        self.assertIn("Nama", template_map_text)
        self.assertIn("Kelas", template_map_text)
        self.assertIn("Fisika", template_map_text)

        loaded_skill = Skill.from_dir(skill_dir)
        tools = {name: fn for name, fn, _ in loaded_skill.tools}
        result = tools[saved["template_tool"]](
            asset_path="assets/form-penilaian-siswa.docx",
            fields={
                "nama": "Budi",
                "kelas": "12A",
                "fisika": 80,
                "kimia": 100,
                "matematika": 90,
                "biologi": 68,
            },
            output_filename="penilaian-budi.pdf",
        )
        self.assertNotIn("error", result)
        self.assertGreaterEqual(result["filledFields"], 6)
        self.assertTrue(result["name"].endswith(".docx"))

        output_doc = Document(Path(result["url"]))
        text = "\n".join(
            cell.text
            for table in output_doc.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("Budi", text)
        self.assertIn("12A", text)
        self.assertIn("80", text)
        self.assertIn("100", text)
        self.assertIn("90", text)
        self.assertIn("68", text)

    def test_docx_template_tool_fills_generic_table_headers(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is required for DOCX template verification")

        upload_dir = Path(self._tmp.name) / "uploads"
        upload_dir.mkdir()
        template = upload_dir / "invoice-template.docx"
        document = Document()
        document.add_heading("INVOICE", level=1)
        document.add_paragraph("Nomor Invoice:")
        document.add_paragraph("Pelanggan:")
        table = document.add_table(rows=3, cols=3)
        for index, header in enumerate(["Produk", "Jumlah", "Harga"]):
            table.cell(0, index).text = header
        document.save(template)

        saved = self.store.save_from_prompt(
            "Buat skill untuk mengisi invoice dari template Word.",
            uploads=[
                {
                    "filename": "invoice-template.docx",
                    "source_name": "invoice-template.docx",
                    "bucket": "assets",
                    "source_path": str(template),
                }
            ],
        )

        from openbench.intelligence.skill import Skill

        skill_dir = self.store.paths()[0]
        loaded_skill = Skill.from_dir(skill_dir)
        tools = {name: fn for name, fn, _ in loaded_skill.tools}
        result = tools[saved["template_tool"]](
            asset_path="assets/invoice-template.docx",
            fields={"Nomor Invoice": "INV-001", "Pelanggan": "PT Contoh"},
            records=[
                {"Produk": "Laptop", "Jumlah": 2, "Harga": "12000000"},
                {"Produk": "Mouse", "Jumlah": 5, "Harga": "150000"},
            ],
            output_filename="invoice-output.docx",
        )
        self.assertNotIn("error", result)

        output_doc = Document(Path(result["url"]))
        text = "\n".join(
            [paragraph.text for paragraph in output_doc.paragraphs]
            + [
                cell.text
                for table in output_doc.tables
                for row in table.rows
                for cell in row.cells
            ]
        )
        self.assertIn("Nomor Invoice: INV-001", text)
        self.assertIn("Pelanggan: PT Contoh", text)
        self.assertIn("Laptop", text)
        self.assertIn("Mouse", text)
        self.assertIn("12000000", text)
        self.assertIn("150000", text)

    def test_save_from_prompt_reuses_matching_mcp_tool_before_creating_function(self):
        class FakeMCPRegistry:
            def list_payload(self):
                return {
                    "servers": [
                        {
                            "id": "image-server",
                            "name": "image",
                            "enabled": True,
                            "tools": [
                                {
                                    "name": "generate_image",
                                    "registered_tool_name": "image_generate_image",
                                    "description": "Generate image previews",
                                    "enabled": True,
                                }
                            ],
                        }
                    ]
                }

        functions = CustomFunctionStore(self._tmp.name)
        saved = self.store.save_from_prompt(
            "Buat skill kampanye yang perlu generate image untuk poster.",
            custom_functions=functions,
            mcp_registry=FakeMCPRegistry(),
        )

        tooling = saved["tooling"]
        self.assertEqual(tooling["created_functions"], [])
        self.assertEqual(tooling["required"][0]["type"], "mcp")
        self.assertEqual(tooling["required"][0]["name"], "image_generate_image")
        self.assertFalse((functions.root / "custom_skill_image_brief.py").exists())

    def test_save_markdown_updates_metadata_from_skill_md(self):
        saved = self.store.save_from_prompt("Buat skill untuk review risiko.")
        updated = self.store.save_markdown(
            saved["id"],
            "\n".join(
                [
                    "# Audit SOP",
                    "",
                    "Membantu audit internal.",
                    "",
                    "## Triggers",
                    "",
                    "- audit internal",
                    "",
                    "## Instructions",
                    "",
                    "Selalu susun temuan dan rekomendasi.",
                    "",
                    "## Version",
                    "",
                    "0.2.0",
                ]
            ),
        )
        self.assertEqual(updated["id"], saved["id"])
        self.assertEqual(updated["name"], "Audit SOP")
        self.assertEqual(updated["version"], "0.2.0")
        self.assertEqual(updated["triggers"], ["audit internal"])
        self.assertIn("Selalu susun", updated["instructions"])


class TestCustomSkillRoutes(unittest.TestCase):
    def _client(self) -> TestClient:
        stack = ExitStack()
        self.addCleanup(stack.close)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        tmpdir = Path(tmp.name)
        stack.enter_context(
            patch.dict(
                environ,
                {
                    "GENERAL_CHAT_STORAGE_ROOT": str(tmpdir / "storage"),
                    "GENERAL_CHAT_UPLOAD_DIR": str(tmpdir / "uploads"),
                    "GENERAL_CHAT_DOWNLOAD_DIR": str(tmpdir / "downloads"),
                    "GENERAL_CHAT_FIREBASE_PROJECT_ID": "demo-project",
                    "GENERAL_CHAT_BOOTSTRAP_ADMIN": "boss@example.com",
                    "OPENBENCH_PROFILE_DIR": str(tmpdir / "profiles"),
                    "GENERAL_CHAT_CUSTOM_SKILLS_DIR": str(tmpdir / "skills"),
                },
                clear=False,
            )
        )
        environ.pop("OPENBENCH_AUTH_DISABLED", None)

        self.create_agent_calls: list[dict] = []

        def _fake_create_agent(**kwargs):
            self.create_agent_calls.append(kwargs)
            agent = Mock()
            agent.model = "mock-model"
            agent._persona = kwargs.get("persona")
            agent._skill_registry = None
            return agent

        stack.enter_context(
            patch("general_chat.server.app.create_agent", side_effect=_fake_create_agent)
        )

        verifier = Mock()
        verifier.verify.return_value = ADMIN
        verifier_cls = stack.enter_context(patch("general_chat.server.auth.FirebaseIDVerifier"))
        verifier_cls.return_value = verifier

        from general_chat.server import auth as auth_module
        from general_chat.server.app import create_app

        auth_module._verifier.cache_clear()
        self.addCleanup(auth_module._verifier.cache_clear)
        return TestClient(create_app())

    def test_routes_require_admin_auth(self):
        client = self._client()
        self.assertEqual(client.get("/admin/custom-skills").status_code, 401)

    def test_save_list_delete_flow_rebuilds_agent_with_skill_path(self):
        client = self._client()
        builds_before = len(self.create_agent_calls)
        response = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            json={
                "prompt": (
                    "Buat skill untuk risk review. Return a table with risk, impact, "
                    "mitigation, and owner."
                ),
            },
        )
        self.assertEqual(response.status_code, 200)
        created_id = response.json()["id"]
        self.assertTrue(created_id)
        self.assertEqual(len(self.create_agent_calls), builds_before + 1)
        loaded_paths = self.create_agent_calls[-1]["custom_skill_paths"]
        self.assertEqual(len(loaded_paths), 1)
        self.assertTrue((Path(loaded_paths[0]) / "SKILL.md").is_file())

        listed = client.get("/admin/custom-skills", headers=ADMIN_H)
        self.assertEqual(listed.status_code, 200)
        self.assertEqual([skill["id"] for skill in listed.json()["skills"]], [created_id])

        edited = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            json={
                "id": created_id,
                "skill_md": "# Risk Review\n\nUpdated.\n\n## Version\n\n0.1.1",
            },
        )
        self.assertEqual(edited.status_code, 200)
        self.assertEqual(edited.json()["name"], "Risk Review")

        deleted = client.delete(f"/admin/custom-skills/{created_id}", headers=ADMIN_H)
        self.assertEqual(deleted.status_code, 200)
        second_delete = client.delete(f"/admin/custom-skills/{created_id}", headers=ADMIN_H)
        self.assertEqual(second_delete.status_code, 404)

    def test_save_invalidates_cached_profile_agents(self):
        client = self._client()
        client.post(
            "/admin/agents",
            headers=ADMIN_H,
            json={"name": "Analis", "description": "Keuangan."},
        )
        registry = client.app.state.agent_registry
        first = registry.get("analis")
        self.assertIsNotNone(first)
        response = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            json={"prompt": "Buat skill untuk audit internal singkat."},
        )
        self.assertEqual(response.status_code, 200)
        # The cached profile agent must be rebuilt to pick up the skill.
        second = registry.get("analis")
        self.assertIsNot(first, second)

    def test_save_validation_errors_are_400(self):
        client = self._client()
        response = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            json={"prompt": ""},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("prompt is required", response.json()["detail"])

    def test_prompt_route_can_create_custom_function_dependency(self):
        client = self._client()
        response = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            json={
                "prompt": (
                    "Buat skill untuk menghitung luas bangunan persegi panjang. Saat user "
                    "memberikan panjang dan lebar tanah, berikan hasil luas yang akurat "
                    "dan konsisten."
                ),
            },
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(
            payload["tooling"]["created_functions"][0]["name"],
            "hitung_luas_bangunan",
        )
        self.assertIn('name="hitung_luas_bangunan"', payload["skill_md"])

    def test_prompt_route_accepts_uploaded_skill_resources(self):
        client = self._client()
        response = client.post(
            "/admin/custom-skills",
            headers=ADMIN_H,
            data={
                "prompt": "Buat skill invoice analyzer dengan SOP validasi invoice.",
                "resource_hint": "SOP sebagai reference.",
            },
            files={"files": ("invoice-sop.md", b"# SOP\n\nCek total invoice.", "text/markdown")},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertGreaterEqual(len(payload["resources"]["references"]), 1)
        skill_path = Path(payload["source"])
        self.assertTrue((skill_path / "references" / "invoice-sop.md").is_file())
        self.assertIn("references/invoice-sop.md", payload["skill_md"])


if __name__ == "__main__":
    unittest.main()
